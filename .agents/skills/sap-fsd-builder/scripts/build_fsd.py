# -*- coding: utf-8 -*-
"""Pembangun FSD RICEF Enterprise di atas template Word resmi (.docx & .md).

Pakai:  python build_fsd.py <content.py> <out.docx> <out.md>
Env (opsional):
  FSD_TPL    path template .docx   (default: templates/fsd_master_template.docx)
  FSD_MEDIA  folder gambar         (default: folder content.py)
  TOCPAGES   json nomor halaman daftar isi (hasil toc_resolver.py) untuk pass ke-2
content.py harus mendefinisikan: META (dict), IMG (dict nomor->nama file), C (list elemen AST).
"""
import copy
import sys
import os
import struct
import json
import importlib.util
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Emu

if len(sys.argv) < 4:
    print("Penggunaan: python build_fsd.py <content.py> <out.docx> <out.md>")
    sys.exit(1)

CONTENT, OUT, MD = sys.argv[1], sys.argv[2], sys.argv[3]
spec = importlib.util.spec_from_file_location('fsd_content', CONTENT)
mod = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mod)
C, IMG, META = mod.C, getattr(mod, 'IMG', {}), mod.META

# Cari default template di folder templates/ relatif terhadap script
DEFAULT_TPL = Path(__file__).resolve().parent.parent / "templates" / "fsd_master_template.docx"
TPL = os.environ.get('FSD_TPL', str(DEFAULT_TPL))
MEDIA = os.environ.get('FSD_MEDIA', os.path.dirname(os.path.abspath(CONTENT)))

if not os.path.exists(TPL):
    raise FileNotFoundError(f"Template master tidak ditemukan: {TPL}")

doc = Document(TPL)
body = doc.element.body
kids = list(body)

# Prototipe elemen template (indeks template resmi bersih)
P_H1 = copy.deepcopy(kids[37])
P_H2 = copy.deepcopy(kids[38])
P_H3 = copy.deepcopy(kids[39])
P_BODY = copy.deepcopy(kids[40])
P_BUL = copy.deepcopy(kids[41])
TBL_PR = copy.deepcopy(kids[42].find(qn('w:tblPr')))
P_CAP = copy.deepcopy(kids[43])
P_EXP = copy.deepcopy(kids[44])
SECT = kids[-1]


def strip_runs(p):
    for ch in list(p):
        if ch.tag != qn('w:pPr'):
            p.remove(ch)


def add_run(p, text, bold=False, sz=None):
    r = OxmlElement('w:r')
    if bold or sz:
        rpr = OxmlElement('w:rPr')
        if bold:
            rpr.append(OxmlElement('w:b'))
        if sz:
            s = OxmlElement('w:sz')
            s.set(qn('w:val'), str(sz))
            rpr.append(s)
        r.append(rpr)
    t = OxmlElement('w:t')
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    p.append(r)


def add_inline(p, text, sz=None, allbold=False):
    for i, part in enumerate(text.split('**')):
        if part:
            add_run(p, part, bold=(i % 2 == 1) or allbold, sz=sz)


def mk(proto, text, sz=None, allbold=False):
    p = copy.deepcopy(proto)
    strip_runs(p)
    add_inline(p, text, sz=sz, allbold=allbold)
    return p


def mk_tbl(headers, rows, widths):
    tbl = OxmlElement('w:tbl')
    tbl.append(copy.deepcopy(TBL_PR))
    grid = OxmlElement('w:tblGrid')
    for w in widths:
        g = OxmlElement('w:gridCol')
        g.set(qn('w:w'), str(w))
        grid.append(g)
    tbl.append(grid)

    def row(cells, header=False):
        tr = OxmlElement('w:tr')
        trpr = OxmlElement('w:trPr')
        trpr.append(OxmlElement('w:tblHeader' if header else 'w:cantSplit'))
        tr.append(trpr)
        for c, w in zip(cells, widths):
            tc = OxmlElement('w:tc')
            tcpr = OxmlElement('w:tcPr')
            tw = OxmlElement('w:tcW')
            tw.set(qn('w:type'), 'dxa')
            tw.set(qn('w:w'), str(w))
            tcpr.append(tw)
            if header:
                sh = OxmlElement('w:shd')
                sh.set(qn('w:val'), 'clear')
                sh.set(qn('w:color'), 'auto')
                sh.set(qn('w:fill'), 'D9E2F3')
                tcpr.append(sh)
            tc.append(tcpr)
            for line in str(c).split('\n'):
                p = OxmlElement('w:p')
                add_inline(p, line, sz=17, allbold=header)
                tc.append(p)
            tr.append(tc)
        return tr

    tbl.append(row(headers, True))
    for r in rows:
        assert len(r) == len(widths), (r, widths)
        tbl.append(row(r))
    return tbl


def png_size(path):
    with open(path, 'rb') as f:
        d = f.read(24)
    return struct.unpack('>II', d[16:24])


def set_text_in_runs(el, old, new):
    for t in el.iter(qn('w:t')):
        if t.text and old in t.text:
            t.text = t.text.replace(old, new)
            return True
    return False


# ---- Halaman Judul (Cover)
# 1. Modul
for t in kids[8].iter(qn('w:t')):
    if t.text and ('Procure to Pay' in t.text or 'MM' in t.text or 'Module' in t.text):
        t.text = META.get('cover_module', t.text)
        break

# 2. Judul Spesifikasi
for t in kids[8].iter(qn('w:t')):
    if t.text and ('Purchase Order' in t.text or 'F-MM-01' in t.text or 'REQ-' in t.text or 'Title' in t.text):
        t.text = META.get('cover_title', t.text)
        break

# 3. Organisasi Klien & Nama Proyek
client_org = META.get('cover_org') or META.get('client_name')
proj_name = META.get('project_name', '')
if client_org:
    combined_org = f"{client_org} / {proj_name}" if proj_name else client_org
    for t in kids[8].iter(qn('w:t')):
        if t.text and any(k in t.text for k in ('Enterprise Client', 'Client', 'Project', 'Rollout', 'Organization')):
            t.text = combined_org
            break

# 4. Tanggal
if 'date' in META:
    for t in kids[24].iter(qn('w:t')):
        if t.text and ('10/10/2026' in t.text or '/' in t.text):
            t.text = META['date']
            break

# ---- Header Setiap Halaman
header_title = META.get('header_new', META.get('cover_title', 'Functional Specification'))
for sec in doc.sections:
    for h in (sec.header, sec.first_page_header, sec.even_page_header):
        for t in h._element.iter(qn('w:t')):
            if t.text and ('Form PO' in t.text or 'F-MM-01' in t.text or 'Header' in t.text):
                t.text = header_title
            elif t.text and ('Procure to Pay' in t.text or 'MM' in t.text):
                t.text = META.get('cover_module', t.text)

# ---- Daftar Isi (Pass ke-2 bila ada TOCPAGES)
TOCP = os.environ.get('TOCPAGES')
if TOCP and os.path.exists(TOCP):
    pages = json.load(open(TOCP))
    tps = list(kids[33].iter(qn('w:p')))[:len(pages)]
    for p, pg in zip(tps, pages):
        ts = [t for t in p.iter(qn('w:t')) if t.text and t.text.strip()]
        if ts and pg is not None:
            ts[-1].text = str(pg)


# ---- Tabel Versi
def set_cell(tc, text):
    ts = list(tc.iter(qn('w:t')))
    if ts:
        ts[0].text = text
        for t in ts[1:]:
            t.text = ''


vt = kids[28]
trs = vt.findall(qn('w:tr'))
for tr in trs[2:]:
    vt.remove(tr)
base = trs[1]
last = base
for n, (ver, dt, nm, why) in enumerate(META.get('versions', [])):
    r = base if n == 0 else copy.deepcopy(base)
    if n:
        last.addnext(r)
        last = r
    cs = r.findall(qn('w:tc'))
    set_cell(cs[0], ver)
    set_cell(cs[1], dt)
    set_cell(cs[2], nm)
    set_cell(cs[3], why)

# ---- Bersihkan Isi Lama
for el in kids[37:-1]:
    body.remove(el)

# ---- Rakit Isi Baru dari AST
new, md = [], []
doc_body = doc._body
for item in C:
    k = item[0]
    if k == 'h1':
        new.append(mk(P_H1, item[1]))
        md.append('\n# %s\n' % item[1])
    elif k == 'h2':
        new.append(mk(P_H2, item[1]))
        md.append('\n## %s\n' % item[1])
    elif k == 'h3':
        new.append(mk(P_H3, item[1]))
        md.append('\n### %s\n' % item[1])
    elif k == 'p':
        new.append(mk(P_BODY, item[1]))
        md.append(item[1] + '\n')
    elif k == 'b':
        new.append(mk(P_BUL, item[1]))
        md.append('* ' + item[1])
    elif k == 'tbl':
        _, hd, rows, wd = item
        new.append(mk_tbl(hd, rows, wd))
        new.append(mk(P_BODY, ''))
        md.append('\n| ' + ' | '.join(hd) + ' |')
        md.append('| ' + ' | '.join('---' for _ in hd) + ' |')
        for r in rows:
            md.append('| ' + ' | '.join(str(c).replace('\n', '<br>').replace('|', '\\|') for c in r) + ' |')
        md.append('')
    elif k == 'fig':
        _, n, cap, exp = item
        img_filename = IMG.get(n)
        if img_filename:
            path = os.path.join(MEDIA, img_filename)
            if os.path.exists(path):
                w, h = png_size(path)
                scale = min(5940000 / w, 5400000 / h, 9525)
                pp = copy.deepcopy(P_BODY)
                strip_runs(pp)
                Paragraph(pp, doc_body).add_run().add_picture(path, width=Emu(int(w * scale)), height=Emu(int(h * scale)))
                new.append(pp)
                new.append(mk(P_CAP, cap, sz=18, allbold=True))
                new.append(mk(P_EXP, exp, sz=18))
                md.append('\n**%s**\n\n%s\n\n_(screenshot: lihat file Word)_\n' % (cap, exp))

for el in new:
    SECT.addprevious(el)

doc.core_properties.title = META.get('doc_title', META['cover_title'])
doc.save(OUT)

head = "# %s\n\n%s\n\nTanggal: %s. Dokumen ini digenerate otomatis dari sap-fsd-builder.\n" % (
    META.get('doc_title', META['cover_title']),
    META['cover_module'],
    META.get('date', '')
)
with open(MD, 'w', encoding='utf-8') as f:
    f.write(head + '\n'.join(md) + '\n')

print(f"[OK] FSD Berhasil dirakit: {OUT} ({len(new)} elemen)")
